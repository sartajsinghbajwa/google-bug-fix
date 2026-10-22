"""Event-scoped guest CRM: careful spreadsheet import, bulk grouping and complete exports."""
from __future__ import annotations

import base64
import binascii
import io
import json
import secrets
import zipfile

from fastapi import HTTPException, Request, Response
from openpyxl import load_workbook
from rsvp_features import workbook
from saas_access import enforce as saas_enforce

COLUMNS = ('name','email','phone','party_size','group_name','category','notes','rsvp')
STATUSES = {'Pending','Confirmed','Maybe','Declined'}
MAX_ROWS = 1000
MAX_UPLOAD_BYTES = 2 * 1024 * 1024


def clean_guest(model, values):
    if not isinstance(values, dict):
        raise HTTPException(422, 'Guest rows must be objects')
    record = {}
    for key in COLUMNS:
        if key == 'party_size':
            raw = values.get(key,1)
            if isinstance(raw,bool):
                raise HTTPException(422, 'Seats must be a number')
            try:
                record[key] = int(raw or 1)
                if str(raw).strip() not in ('','1') and float(raw) != record[key]:
                    raise ValueError()
            except (ValueError, TypeError):
                raise HTTPException(422, 'Seats must be a whole number')
        else:
            record[key] = str(values.get(key,'') or '').strip()
    record['category'] = record['category'] or 'General'
    record['rsvp'] = record['rsvp'] or 'Pending'
    if record['rsvp'] not in STATUSES:
        raise HTTPException(422, 'Unsupported RSVP value')
    try:
        return model.model_validate(record)
    except Exception:
        raise HTTPException(422, 'Invalid guest row: check names, seat counts and field lengths')


def identifiers(guest):
    if guest.email.strip():
        return [('email',guest.email.strip().casefold())]
    if guest.phone.strip():
        return [('phone',''.join(ch for ch in guest.phone if ch.isdigit()) or guest.phone.strip().casefold())]
    return [('name_group',guest.name.strip().casefold(),guest.group_name.strip().casefold())]


def install_guest_crm(app, db, access, check_event, guest_model, uid, now):
    def store_rows(con, wid, eid, rows):
        if not isinstance(rows,list) or len(rows)>MAX_ROWS:
            raise HTTPException(422, 'Import at most 1,000 guests at a time')
        people=[clean_guest(guest_model,item) for item in rows]
        saved=con.execute('''SELECT name,email,phone,group_name FROM guests
                            WHERE workspace_id=? AND event_id=?''',(wid,eid)).fetchall()
        known=set()
        for row in saved:
            fake=type('GuestKey', (), dict(row))()
            known.update(identifiers(fake))
        imported=skipped=0
        for guest in people:
            keys=identifiers(guest)
            if any(k in known for k in keys):
                skipped+=1
                continue
            saas_enforce(con,wid,'guests')
            con.execute('''INSERT INTO guests
                (id,workspace_id,event_id,name,email,phone,party_size,group_name,category,notes,
                 rsvp,invitation_token,created_at)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)''',
                (uid(),wid,eid,guest.name.strip(),guest.email.strip(),guest.phone.strip(),
                 guest.party_size,guest.group_name.strip(),guest.category.strip(),guest.notes.strip(),
                 guest.rsvp,secrets.token_urlsafe(32),now()))
            known.update(keys)
            imported+=1
        return {'imported':imported,'duplicates_skipped':skipped,'total':len(people)}

    @app.post('/api/w/{wid}/events/{eid}/guests/import-reviewed')
    def import_reviewed(request:Request,wid:str,eid:str,payload:dict):
        if set(payload)!={'rows'}:
            raise HTTPException(422,'Expected rows list')
        with db() as con:
            access(request,con,wid,True)
            check_event(con,wid,eid)
            return store_rows(con,wid,eid,payload['rows'])

    @app.post('/api/w/{wid}/events/{eid}/guests/import.xlsx')
    def import_excel(request:Request,wid:str,eid:str,payload:dict):
        with db() as con:
            access(request,con,wid,True)
            check_event(con,wid,eid)
            if set(payload)!={'data_base64'} or not isinstance(payload.get('data_base64'),str):
                raise HTTPException(422,'Upload an Excel workbook')
            encoded=payload['data_base64']
            if len(encoded)>((MAX_UPLOAD_BYTES+2)//3*4)+4:
                raise HTTPException(413,'Excel workbook exceeds 2 MB')
            try:
                data=base64.b64decode(encoded,validate=True)
            except (ValueError,binascii.Error):
                raise HTTPException(422,'Invalid Excel file')
            if not data or len(data)>MAX_UPLOAD_BYTES:
                raise HTTPException(413,'Excel workbook exceeds 2 MB')
            try:
                with zipfile.ZipFile(io.BytesIO(data)) as archive:
                    files=archive.infolist()
                    if len(files)>100 or sum(x.file_size for x in files)>12*1024*1024:
                        raise HTTPException(413,'Excel workbook is too large to process')
                wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
                try:
                    sheet=wb.active
                    iterator=sheet.iter_rows(values_only=True,max_row=MAX_ROWS+2)
                    header=next(iterator,None)
                    if not header:
                        raise HTTPException(422,'Excel worksheet is empty')
                    aliases={'reserved_seats':'party_size','group':'group_name',
                             'private_notes':'notes','guest_name':'name'}
                    keys=[aliases.get(str(x or '').strip().lower().replace(' ','_'),
                                      str(x or '').strip().lower().replace(' ','_'))
                          for x in header]
                    if 'name' not in keys or len(set(keys))!=len(keys):
                        raise HTTPException(422,'Excel worksheet needs unique headers including name')
                    rows=[]
                    for cells in iterator:
                        if not any(v is not None and str(v).strip() for v in cells):
                            continue
                        if len(rows)>=MAX_ROWS:
                            raise HTTPException(422,'Import at most 1,000 guests at a time')
                        rows.append({key:val for key,val in zip(keys,cells) if key in COLUMNS})
                finally:
                    wb.close()
            except (OSError,ValueError,zipfile.BadZipFile,KeyError,TypeError) as exc:
                raise HTTPException(422,'Could not read Excel workbook') from exc
            return store_rows(con,wid,eid,rows)

    @app.patch('/api/w/{wid}/events/{eid}/guests/bulk')
    def bulk_group(request:Request,wid:str,eid:str,payload:dict):
        if set(payload) != {'ids','changes'}:
            raise HTTPException(422,'Provide guest IDs and changes')
        ids=payload['ids']
        changes=payload['changes']
        if not isinstance(ids,list) or not 1<=len(ids)<=300 or any(not isinstance(x,str) or len(x)>100 for x in ids) or len(set(ids))!=len(ids):
            raise HTTPException(422,'Select 1–300 unique guests')
        if not isinstance(changes,dict) or not changes or set(changes)-{'group_name','category','rsvp'}:
            raise HTTPException(422,'Only guest group, category and RSVP may be changed together')
        for key,value in changes.items():
            if not isinstance(value,str) or len(value)>(160 if key=='group_name' else 80):
                raise HTTPException(422,'Invalid bulk guest change')
            if key=='rsvp' and value not in STATUSES:
                raise HTTPException(422,'Unsupported RSVP status')
        with db() as con:
            access(request,con,wid,True)
            check_event(con,wid,eid)
            count=con.execute('SELECT COUNT(*) FROM guests WHERE workspace_id=? AND event_id=? AND id IN ('+
                              ','.join('?' for _ in ids)+')',(wid,eid,*ids)).fetchone()[0]
            if count!=len(ids):
                raise HTTPException(404,'Some guests are not in this event')
            sql='UPDATE guests SET '+','.join(key+'=?' for key in changes)+' WHERE workspace_id=? AND event_id=? AND id IN ('+','.join('?' for _ in ids)+')'
            con.execute(sql,(*changes.values(),wid,eid,*ids))
            return {'updated':len(ids)}

    @app.get('/api/w/{wid}/events/{eid}/guests/export.xlsx')
    def export_excel(request:Request,wid:str,eid:str):
        with db() as con:
            access(request,con,wid)
            check_event(con,wid,eid)
            rows=con.execute('''SELECT g.*, r.party_json, r.submitted_at
                FROM guests g LEFT JOIN rsvp_responses r ON r.guest_id=g.id
                WHERE g.workspace_id=? AND g.event_id=? ORDER BY g.group_name,g.name''',(wid,eid)).fetchall()
            headers=['Name','Email','Phone','Reserved seats','Group','Category','Private notes','RSVP',
                     'Attending','Adults','Kids','Plus one','Attendee names','Meal preferences','RSVP submitted']
            output=[]
            for guest in rows:
                try: party=json.loads(guest['party_json'] or '{}')
                except (ValueError,TypeError): party={}
                people=party.get('guests',[]) if isinstance(party.get('guests',[]),list) else []
                output.append(dict(zip(headers,[
                    guest['name'],guest['email'],guest['phone'],guest['party_size'],
                    guest['group_name'],guest['category'],guest['notes'],guest['rsvp'],
                    party.get('total',''),party.get('adults',''),party.get('kids',''),
                    'Yes' if party.get('plus_one') else 'No',
                    ', '.join(x.get('name','') for x in people if isinstance(x,dict) and x.get('name')),
                    ', '.join(x.get('meal','') for x in people if isinstance(x,dict) and x.get('meal')),
                    guest['submitted_at'] or ''
                ])))
        return Response(workbook(output,headers),
            media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition':'attachment; filename="eventos-all-guests.xlsx"',
                     'Cache-Control':'no-store'})
